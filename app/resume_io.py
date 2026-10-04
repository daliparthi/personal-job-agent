"""
Resume files in and out.

Reading (DOCX / PDF / TXT / MD) produces a line-level structure that master.draft() turns into the
master_resume.yaml schema:
{
  "header": ["Jane Doe", "Austin, TX | jane@example.com | 555-0100"],
  "sections": [
    {"title": "PROFESSIONAL SUMMARY", "kind": "summary", "blocks": [{"type": "text", "text": "..."}]},
    {"title": "EXPERIENCE", "kind": "experience", "blocks": [{"type": "text", "text": "Data Engineer, Acme | 2020 - Present"},
                                                            {"type": "bullet", "text": "Built ..."}]}
  ]
}

Writing (text / HTML for PDF / DOCX) renders the master_resume.yaml schema (see master.py).
"""
import html
import io
import re
from datetime import date

SECTION_WORDS = {
    "summary": ["summary", "professional summary", "profile", "professional profile", "objective",
                "career objective", "about me", "career summary", "executive summary", "overview"],
    "skills": ["skills", "technical skills", "core competencies", "competencies", "key skills", "technologies",
               "tools", "skills & tools", "skills and tools", "areas of expertise", "expertise", "tech stack",
               "technical proficiencies", "core skills", "skills summary"],
    "experience": ["experience", "work experience", "professional experience", "employment history",
                   "work history", "career history", "relevant experience", "employment"],
    "projects": ["projects", "key projects", "selected projects", "personal projects", "academic projects"],
    "education": ["education", "education & training", "academic background", "education and training"],
    "certifications": ["certifications", "certificates", "licenses", "licenses & certifications",
                       "certifications & licenses", "training", "certifications and training"],
    "other": ["awards", "honors", "publications", "volunteer", "volunteering", "languages", "interests",
              "achievements", "accomplishments", "activities", "patents", "affiliations", "references",
              "additional information", "leadership"],
}
_HEADINGS = {w: kind for kind, words in SECTION_WORDS.items() for w in words}
_BULLET = re.compile(r"^\s*(?:[•●▪◦‣∙·■□➢➤►\-–\*]|o\s||)\s*")


def _clean(s):
    return re.sub(r"\s+", " ", (s or "").replace("\xa0", " ")).strip()


def _heading_kind(line: str):
    t = re.sub(r"[:#*_]+$|^[#*_]+", "", line.strip()).strip()
    low = t.lower().replace("&amp;", "&")
    if low in _HEADINGS:
        return _HEADINGS[low], t
    words = t.split()
    if (1 <= len(words) <= 5 and t.isupper() and not re.search(r"[|@\d]", t) and len(t) > 3):
        kind = next((k for w, k in _HEADINGS.items() if w in low), "other")
        return kind, t
    return None


def _lines_from_docx(data: bytes):
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(data))
    out = []

    def para(p):
        text = _clean(p.text)
        if not text:
            return
        style = (p.style.name or "").lower() if p.style is not None else ""
        numbered = p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
        is_bullet = "list" in style or numbered or bool(_BULLET.match(p.text))
        is_heading = style.startswith("heading") or style == "title"
        out.append((text, is_bullet, is_heading))

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            para(Paragraph(child, doc))
        elif child.tag == qn("w:tbl"):
            from docx.table import Table
            for row in Table(child, doc).rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        para(p)
    return out


def _lines_from_pdf(data: bytes):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return [(_clean(l), bool(_BULLET.match(l)), False) for l in text.splitlines() if _clean(l)]


def _lines_from_text(data: bytes):
    text = data.decode("utf-8", errors="replace")
    out = []
    for l in text.splitlines():
        if not l.strip():
            out.append(("", False, False))
            continue
        is_md_heading = bool(re.match(r"^\s*#{1,6}\s", l))
        out.append((_clean(re.sub(r"^\s*#{1,6}\s", "", l)), bool(_BULLET.match(l)), is_md_heading))
    return out


def parse_resume(filename: str, data: bytes):
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "docx":
        lines, wrapped = _lines_from_docx(data), False
    elif ext == "pdf":
        lines, wrapped = _lines_from_pdf(data), True
    elif ext in ("txt", "md", "text", "markdown"):
        lines, wrapped = _lines_from_text(data), True
    else:
        raise ValueError("Upload a .docx, .pdf, .txt or .md resume")
    return structure_from_lines(lines, wrapped)


def structure_from_lines(lines, wrapped_lines: bool):
    header, sections, cur = [], [], None
    blank_before = False
    last_line_len = 0

    def title_like(s):
        words = re.findall(r"[A-Za-z][\w&/.'-]*", s)
        small = {"of", "and", "the", "in", "for", "at", "to", "a", "an", "&"}
        caps = sum(1 for w in words if w[0].isupper() or w.lower() in small)
        return bool(words) and len(words) <= 10 and caps / len(words) >= 0.7

    for text, is_bullet, is_heading in lines:
        if not text:
            blank_before = True
            continue
        hk = _heading_kind(text)
        if hk and cur is None and hk[0] == "other" and text.lower().strip(" :") not in _HEADINGS:
            hk = None  # an all-caps line above the first real section is the name / headline, not a heading
        if hk or (is_heading and cur is not None and len(text.split()) <= 6):
            kind, title = hk if hk else ("other", text)
            cur = {"title": title, "kind": kind, "blocks": []}
            sections.append(cur)
            blank_before = False
            continue
        if cur is None:
            header.append(text)
            continue
        body = _BULLET.sub("", text) if is_bullet else text
        prev = cur["blocks"][-1] if cur["blocks"] else None
        # PDFs and text files wrap long lines; glue continuation lines back on.
        continuation = (wrapped_lines and prev is not None and not is_bullet and not blank_before
                        and (prev["type"] == "bullet" or cur["kind"] == "summary")
                        and (body[:1].islower() or (last_line_len >= 70 and not title_like(body)
                                                    and not re.search(r"[.!?]$", prev["text"])))
                        and not re.search(r"\b(19|20)\d{2}\b", body))
        if continuation:
            prev["text"] = f"{prev['text']} {body}"
        elif cur["kind"] == "summary" and prev is not None and not blank_before and not is_bullet and wrapped_lines:
            prev["text"] = f"{prev['text']} {body}"
        else:
            # PDFs often drop bullet glyphs: in experience sections a sentence-like line is a bullet.
            inferred = (wrapped_lines and not is_bullet and cur["kind"] in ("experience", "projects")
                        and len(body) > 40 and not title_like(body) and not re.search(r"\b(19|20)\d{2}\b", body))
            cur["blocks"].append({"type": "bullet" if is_bullet or inferred else "text", "text": body})
        blank_before = False
        last_line_len = len(text)
    if not sections:  # nothing recognisable: treat everything after the first two lines as one section
        sections = [{"title": "EXPERIENCE", "kind": "experience",
                     "blocks": [{"type": "text", "text": h} for h in header[2:]]}]
        header = header[:2]
    return {"header": header, "sections": sections}




# ---------------------------------------------------------------- rendering (master_resume.yaml schema)
ENTRY_KINDS = ("experience", "projects", "education")
ENTRY_FIELDS = {  # the first two make the bold part of an entry's heading line
    "experience": ("title", "company", "location", "start", "end"),
    "education": ("degree", "school", "location", "start", "end"),
    "projects": ("name", "organization", "location", "start", "end"),
}


def split_items(text: str):
    """Split "a, b (c, d); e" on separators that are not inside brackets."""
    out, cur, depth = [], "", 0
    for ch in text:
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth = max(0, depth - 1)
        if depth == 0 and ch in ",;|•·":
            out.append(cur)
            cur = ""
        else:
            cur += ch
    out.append(cur)
    return [x.strip() for x in out if x.strip()]


def group_text(g: dict) -> str:
    items = ", ".join(g.get("items") or [])
    return f"{g['name']}: {items}" if g.get("name") else items


def entry_dates(e: dict) -> str:
    start, end = e.get("start") or "", e.get("end") or ""
    return f"{start} – {end}" if start and end else start or end


def entry_heading(kind: str, e: dict):
    """[(bold part, rest)] lines for an entry's heading."""
    if e.get("heading"):
        return [(line, "") for line in e["heading"]]
    f = ENTRY_FIELDS[kind]
    main = " — ".join(x for x in (e.get(f[0]), e.get(f[1])) if x)
    meta = " | ".join(x for x in (e.get("location"), entry_dates(e)) if x)
    return [(main, meta)] if main or meta else []


def clean_resume(res: dict) -> dict:
    """Drop the empty/None placeholders a tailored copy can hold (e.g. an undone 'Additional' skills line)."""
    out = dict(res)
    secs = []
    for s in res.get("sections") or []:
        s = dict(s)
        for key in ("groups", "bullets", "items", "lines", "entries"):
            if key in s:
                s[key] = [x for x in s[key] or [] if x]
        secs.append(s)
    out["sections"] = secs
    return out


def layout(res: dict):
    """Flat list of printable items: (kind, text[, rest]) with kind in name/contact/h2/p/bullet/role."""
    res = clean_resume(res or {})
    items = []
    if res.get("name"):
        items.append(("name", res["name"]))
    if res.get("headline"):
        items.append(("contact", res["headline"]))
    c = res.get("contact") or {}
    line = " | ".join(x for x in [c.get("location"), c.get("phone"), c.get("email"), *(c.get("links") or [])] if x)
    if line:
        items.append(("contact", line))
    items += [("contact", o) for o in c.get("other") or []]
    for s in res.get("sections") or []:
        kind, body = s.get("kind"), []
        if kind == "summary":
            body += [("p", t) for t in (s.get("text") or "").split("\n") if t.strip()]
            body += [("bullet", b) for b in s.get("bullets") or []]
        elif kind == "skills":
            body += [("p", group_text(g)) for g in s.get("groups") or []]
            body += [("p", t) for t in s.get("lines") or []]
        elif kind in ENTRY_KINDS:
            for e in s.get("entries") or []:
                body += [("role", main, meta) for main, meta in entry_heading(kind, e)]
                if e.get("description"):
                    body.append(("p", e["description"]))
                body += [("bullet", b) for b in e.get("bullets") or []]
        else:
            style = "bullet" if s.get("style") == "bullets" else "p"
            body += [(style, t) for t in s.get("items") or []]
        if body:
            items.append(("h2", (s.get("title") or kind or "").upper()))
            items += body
    return items


def to_text(res: dict) -> str:
    out = []
    for it in layout(res):
        kind = it[0]
        if kind == "h2":
            out += ["", it[1]]
        elif kind == "bullet":
            out.append(f"• {it[1]}")
        elif kind == "role":
            out.append(" | ".join(x for x in it[1:] if x))
        else:
            out.append(it[1])
    return "\n".join(out).strip() + "\n"


def to_html(res: dict, title="Resume") -> str:
    parts = []
    for it in layout(res):
        kind, text = it[0], html.escape(it[1])
        if kind == "name":
            parts.append(f"<h1>{text}</h1>")
        elif kind == "contact":
            parts.append(f'<p class="contact">{text}</p>')
        elif kind == "h2":
            parts.append(f"<h2>{text}</h2>")
        elif kind == "bullet":
            # A real "•" character (not a CSS list marker) so ATS text extraction sees each bullet.
            parts.append(f'<p class="bullet">• {text}</p>')
        elif kind == "role":
            rest = f" | {html.escape(it[2])}" if it[2] and it[1] else html.escape(it[2])
            parts.append(f'<p class="role"><b>{text}</b>{rest}</p>')
        else:
            parts.append(f"<p>{text}</p>")
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
 body {{ font-family: Calibri, Arial, Helvetica, sans-serif; font-size: 10.5pt; color: #111; margin: 0; line-height: 1.32;
        /* ligatures turn "fi"/"fl" into single glyphs that ATS text extraction reads as garbage */
        font-variant-ligatures: none; font-feature-settings: "liga" 0, "clig" 0, "dlig" 0; }}
 h1 {{ font-size: 18pt; margin: 0 0 2pt; text-align: center; }}
 .contact {{ text-align: center; margin: 0; font-size: 9.5pt; }}
 h2 {{ font-size: 11pt; letter-spacing: .04em; border-bottom: 1px solid #444; margin: 12pt 0 4pt; padding-bottom: 1pt; }}
 p {{ margin: 0 0 3pt; }}
 p.role {{ margin-top: 6pt; }}
 p.bullet {{ padding-left: 12pt; text-indent: -9pt; margin: 0 0 2pt 4pt; }}
</style></head><body>{''.join(parts)}</body></html>"""


def to_docx(res: dict, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.7)
        sec.top_margin = sec.bottom_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    for it in layout(res):
        kind, text = it[0], it[1]
        if kind == "name":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(16)
        elif kind == "contact":
            doc.add_paragraph(text).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(11)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        elif kind == "role":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.add_run(text).bold = True
            if it[2]:
                p.add_run(f" | {it[2]}" if text else it[2])
        else:
            doc.add_paragraph(text)
    doc.save(path)


# ---------------------------------------------------------------- cover letter (written next to the resume)
def clean_letter(letter):
    """The cover letter the page saved ({greeting, paragraphs, signoff, name, date, answers, changes}) reduced to
    the text that is written out, or None when there is no letter."""
    if not isinstance(letter, dict):
        return None
    paragraphs = [p.strip() for p in letter.get("paragraphs") or [] if isinstance(p, str) and p.strip()]
    if not paragraphs:
        return None
    answers = []
    for a in letter.get("answers") or []:
        if isinstance(a, dict) and str(a.get("text") or "").strip():
            answers.append({"question": str(a.get("question") or "").strip(), "text": str(a["text"]).strip()})
    return {"greeting": str(letter.get("greeting") or "").strip(), "paragraphs": paragraphs,
            "signoff": str(letter.get("signoff") or "Sincerely,").strip(), "name": str(letter.get("name") or "").strip(),
            "date": str(letter.get("date") or "").strip(), "answers": answers}


def long_date(iso: str) -> str:
    """'2026-10-04' -> 'October 4, 2026' (anything else is kept as it is)."""
    try:
        d = date.fromisoformat(iso)
    except (TypeError, ValueError):
        return iso or ""
    return f"{d:%B} {d.day}, {d.year}"


def letter_layout(letter: dict, res: dict):
    """(kind, text) items: the resume's name and contact lines, then date, greeting, paragraphs and signature."""
    items = [it for it in layout(res) if it[0] in ("name", "contact")]
    if letter.get("date"):
        items.append(("date", long_date(letter["date"])))
    if letter.get("greeting"):
        items.append(("greeting", letter["greeting"]))
    items += [("p", t) for t in letter["paragraphs"]]
    items.append(("signoff", letter.get("signoff") or "Sincerely,"))
    name = letter.get("name") or (res or {}).get("name")
    if name:
        items.append(("signature", name))
    return items


def letter_text(letter: dict, res: dict) -> str:
    out, prev = [], None
    for kind, text in letter_layout(letter, res):
        if prev and not (prev in ("name", "contact") and kind == "contact") and not (prev == "signoff"):
            out.append("")
        out.append(text)
        prev = kind
    return "\n".join(out).strip() + "\n"


def letter_html(letter: dict, res: dict, title="Cover letter") -> str:
    parts = []
    for kind, text in letter_layout(letter, res):
        t = html.escape(text)
        if kind == "name":
            parts.append(f"<h1>{t}</h1>")
        elif kind == "contact":
            parts.append(f'<p class="contact">{t}</p>')
        elif kind == "signoff":
            parts.append(f'<p class="signoff">{t}</p>')
        elif kind == "signature":
            parts.append(f'<p class="signature">{t}</p>')
        else:
            parts.append(f'<p class="{kind}">{t}</p>')
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
 body {{ font-family: Calibri, Arial, Helvetica, sans-serif; font-size: 11pt; color: #111; margin: 0; line-height: 1.4;
        font-variant-ligatures: none; font-feature-settings: "liga" 0, "clig" 0, "dlig" 0; }}
 h1 {{ font-size: 18pt; margin: 0 0 2pt; text-align: center; }}
 .contact {{ text-align: center; margin: 0; font-size: 9.5pt; }}
 .date {{ margin: 22pt 0 14pt; }}
 p {{ margin: 0 0 10pt; }}
 .signoff {{ margin: 16pt 0 2pt; }}
</style></head><body>{''.join(parts)}</body></html>"""


def letter_docx(letter: dict, res: dict, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for kind, text in letter_layout(letter, res):
        if kind == "name":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(16)
            p.paragraph_format.space_after = Pt(0)
        elif kind == "contact":
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
        elif kind == "date":
            p = doc.add_paragraph(text)
            p.paragraph_format.space_before = Pt(18)
        elif kind == "signoff":
            p = doc.add_paragraph(text)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(0)
        else:
            doc.add_paragraph(text)
    doc.save(path)


def answers_text(letter: dict, job: dict) -> str:
    """The short answers, ready to paste into an application form."""
    out = [f"Short answers for {job.get('title')} at {job.get('company')}", ""]
    for a in letter.get("answers") or []:
        out += [a["question"], a["text"], ""]
    return "\n".join(out).strip() + "\n"

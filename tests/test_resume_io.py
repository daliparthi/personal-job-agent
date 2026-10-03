from pathlib import Path

import pytest
from docx import Document

from app import master, resume_io
from tests.conftest import SAMPLE_MASTER, SAMPLE_RESUME

GOLDEN = Path(__file__).parent / "golden" / "sample_master_resume.yaml"


def test_sample_resume_parses_to_golden_yaml():
    """The rule-based parse of samples/sample_master_resume.txt. If you change the parser on purpose, regenerate:
    python -c "from tests.test_resume_io import regenerate; regenerate()" """
    assert _parse_sample() == GOLDEN.read_text(encoding="utf-8")


def regenerate():
    GOLDEN.write_text(_parse_sample(), encoding="utf-8", newline="\n")


def _parse_sample():
    structured = resume_io.parse_resume(SAMPLE_RESUME.name, SAMPLE_RESUME.read_bytes())
    return master.dump(master.draft(structured))


def test_unsupported_extension():
    with pytest.raises(ValueError, match="Upload a"):
        resume_io.parse_resume("resume.rtf", b"x")


def test_markdown_headings_and_wrapped_lines():
    text = ("# Jane Doe\njane@example.com\n\n## Experience\nEngineer, Acme | 2020 - 2022\n"
            "- Built a very long thing that wraps onto the next line because the PDF export broke it\n"
            "  in two places\n")
    s = resume_io.parse_resume("cv.md", text.encode())
    assert s["header"] == ["Jane Doe", "jane@example.com"]
    exp = s["sections"][0]
    assert exp["kind"] == "experience"
    assert exp["blocks"][-1] == {"type": "bullet", "text": "Built a very long thing that wraps onto the next line "
                                                            "because the PDF export broke it in two places"}


def test_nothing_recognisable_becomes_one_section():
    s = resume_io.parse_resume("cv.txt", b"Jane\nline two\nline three\nline four")
    assert s["header"] == ["Jane", "line two"]
    assert [b["text"] for b in s["sections"][0]["blocks"]] == ["line three", "line four"]


def test_docx_round_trip(tmp_path):
    path = tmp_path / "cv.docx"
    resume_io.to_docx(SAMPLE_MASTER, path)
    texts = [p.text for p in Document(path).paragraphs]
    assert texts[0] == "Jordan Avery"
    assert "SKILLS" in texts and "Languages: Python, SQL" in texts
    s = resume_io.parse_resume("cv.docx", path.read_bytes())
    kinds = [x["kind"] for x in s["sections"]]
    assert kinds == ["summary", "skills", "experience", "education"]
    bullets = [b for b in s["sections"][2]["blocks"] if b["type"] == "bullet"]
    assert len(bullets) == 2


def test_text_and_html_rendering():
    text = resume_io.to_text(SAMPLE_MASTER)
    assert text.startswith("Jordan Avery\n")
    assert "• Designed Spark and Airflow pipelines" in text
    assert "Senior Data Engineer — Northwind Analytics | Austin, TX | 2021 – Present" in text
    html = resume_io.to_html({**SAMPLE_MASTER, "name": "<Jordan>"}, title="T")
    assert "<h1>&lt;Jordan&gt;</h1>" in html and "font-variant-ligatures: none" in html


def test_clean_resume_drops_empty_placeholders():
    res = {"sections": [{"kind": "skills", "groups": [None, {"name": "", "items": ["x"]}], "lines": [""]}]}
    out = resume_io.clean_resume(res)
    assert out["sections"][0]["groups"] == [{"name": "", "items": ["x"]}]
    assert out["sections"][0]["lines"] == []


@pytest.mark.parametrize("text, want", [
    ("a, b (c, d); e", ["a", "b (c, d)", "e"]),
    ("Python | SQL • Spark", ["Python", "SQL", "Spark"]),
    ("", []),
])
def test_split_items(text, want):
    assert resume_io.split_items(text) == want

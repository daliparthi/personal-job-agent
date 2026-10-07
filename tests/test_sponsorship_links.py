"""Sponsorship / H-1B / H-4 EAD signals of a posting, and clickable contact links in the resume files."""
import pytest
from docx import Document

from app import db, master, resume_io, search
from app.jobparse import sponsorship_in
from tests.conftest import SAMPLE_MASTER, make_job


@pytest.mark.parametrize("text, want", [
    ("We offer visa sponsorship for this role.", ["sponsors"]),
    ("Sponsorship is available for qualified candidates.", ["sponsors"]),
    ("We can sponsor H-1B visas.", ["sponsors", "h1b"]),
    ("H-4 EAD holders are welcome to apply.", ["h4ead"]),
    ("We accept candidates on H4 EAD or OPT.", ["h4ead"]),
    ("We are unable to provide visa sponsorship for this role.", ["no"]),
    ("We do not sponsor H-1B visas. H-4 EAD candidates are welcome.", ["h4ead", "no"]),
    ("Python and SQL experience required.", []),
    ("", []),
])
def test_sponsorship_tags(text, want):
    assert sponsorship_in(text) == want


def test_list_filters_by_sponsorship():
    db.upsert_jobs([make_job("a:yes", description_text="We sponsor H-1B visas."),
                    make_job("a:ead", description_text="H-4 EAD is accepted."),
                    make_job("a:no", description_text="We cannot sponsor visas."),
                    make_job("a:silent", description_text="Python and SQL.")])
    search.rescore_all()
    assert db.get_job("a:yes")["sponsorship"] == "sponsors,h1b"
    s = db.get_settings()

    def ids(mode):
        s["filters"]["sponsorship"] = mode
        return {j["id"] for j in search.list_jobs(s)}

    assert ids("any") == {"a:yes", "a:ead", "a:no", "a:silent"}
    assert ids("offered") == {"a:yes"}
    assert ids("h4ead") == {"a:ead"}
    assert ids("not_denied") == {"a:yes", "a:ead", "a:silent"}


@pytest.mark.parametrize("link, want", [
    ("linkedin.com/in/jane", ("LinkedIn", "https://linkedin.com/in/jane")),
    ("https://www.linkedin.com/in/jane/", ("LinkedIn", "https://www.linkedin.com/in/jane/")),
    ("github.com/jane", ("GitHub", "https://github.com/jane")),
    ("jane.dev", ("jane.dev", "https://jane.dev")),
    ({"text": "My site", "url": "jane.dev"}, ("My site", "https://jane.dev")),
    ({"text": "In", "url": "https://linkedin.com/in/jane"}, ("In", "https://linkedin.com/in/jane")),
])
def test_link_parts(link, want):
    assert resume_io.link_parts(link) == want


def test_master_keeps_link_mappings():
    data = master.normalize({"name": "J", "contact": {"links": ["linkedin.com/in/j", {"text": "Site", "url": "j.dev"},
                                                                 {"url": "gh.com/j"}, {"text": "", "url": ""}]},
                             "sections": []})
    assert data["contact"]["links"] == ["linkedin.com/in/j", {"text": "Site", "url": "j.dev"}, "gh.com/j"]
    again = master.parse_yaml(master.dump(data))
    assert again["contact"]["links"] == data["contact"]["links"]


def _with_links(*links):
    return {**SAMPLE_MASTER, "contact": {**SAMPLE_MASTER["contact"], "links": list(links)}}


def test_html_and_text_show_the_short_label():
    res = _with_links("linkedin.com/in/jordan-avery-example")
    html = resume_io.to_html(res)
    assert '<a href="https://linkedin.com/in/jordan-avery-example">LinkedIn</a>' in html
    text = resume_io.to_text(res)
    assert "| LinkedIn" in text and "linkedin.com/in/" not in text


def test_only_headings_and_skill_group_names_are_bold(tmp_path):
    import re
    assert re.findall(r"<b>(.*?)</b>", resume_io.to_html(SAMPLE_MASTER)) == ["Languages:", "Data:"]
    path = tmp_path / "cv.docx"
    resume_io.to_docx(SAMPLE_MASTER, path)
    bold = {p.text for p in Document(path).paragraphs if any(r.bold for r in p.runs)}
    assert "EDUCATION" in bold and "Jordan Avery" in bold
    assert not any("University" in t or "Senior Data Engineer" in t for t in bold)
    skills = next(p for p in Document(path).paragraphs if p.text.startswith("Languages:"))
    assert [r.text for r in skills.runs if r.bold] == ["Languages:"]  # the items after the name are not bold


def test_docx_link_is_a_real_hyperlink(tmp_path):
    path = tmp_path / "cv.docx"
    resume_io.to_docx(_with_links("linkedin.com/in/jordan-avery-example"), path)
    doc = Document(path)
    contact = next(p for p in doc.paragraphs if "LinkedIn" in p.text)
    assert contact.text.endswith("| LinkedIn")
    targets = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert targets == ["https://linkedin.com/in/jordan-avery-example"]
    assert contact._p.xpath("./w:hyperlink")
